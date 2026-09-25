"""Bounded project document helpers and a discoverable, isolated authoring runtime."""
from __future__ import annotations

import hashlib
import io
import json
import math
import os
from pathlib import Path
import stat
import sys
import zipfile
from uuid import uuid4

MAX_FILE = 20 * 1024 * 1024


def runtime_path():
    root = Path(__file__).resolve().parent.parent
    bundled = root / "document_packages"
    return bundled if bundled.is_dir() else root / "dist" / f"document-runtime-{sys.version_info.major}.{sys.version_info.minor}"


def prepare_imports():
    path = runtime_path()
    if not (path / "proto-mind-document-runtime.json").is_file():
        raise ValueError("Document runtime is not installed. Build/install the current Proto-Mind package first.")
    if str(path) not in sys.path: sys.path.insert(0, str(path))
    os.environ["OPENPYXL_DEFUSEDXML"] = "True"
    return path


def environment():
    path = prepare_imports()
    import importlib.metadata
    return {"python": sys.executable, "pythonpath": str(path), "packages": {
        name: importlib.metadata.version(name) for name in ["python-docx", "openpyxl", "python-pptx", "reportlab", "Pillow"]},
        "notice": "Use this Python and PYTHONPATH for rich document authoring. Reopen outputs and inspect rendered pages; a saved file is not visual verification."}


def read_bytes(reader, path):
    parts = reader._relative(path)
    with reader._directory(parts[:-1]) as directory:
        fd = os.open(parts[-1], os.O_RDONLY | os.O_NOFOLLOW | os.O_NONBLOCK, dir_fd=directory)
        try:
            info = os.fstat(fd)
            if not stat.S_ISREG(info.st_mode) or info.st_size > MAX_FILE: raise ValueError("Choose a regular document below 20 MB.")
            with os.fdopen(fd, "rb", closefd=False) as source: data = source.read(MAX_FILE + 1)
            if len(data) > MAX_FILE: raise ValueError("Document grew past its limit.")
            return data
        finally: os.close(fd)


def check_archive(data):
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        entries = archive.infolist()
        if len(entries) > 2000 or sum(x.file_size for x in entries) > 40 * 1024 * 1024:
            raise ValueError("Document archive exceeds its inspection limits.")
        for item in entries:
            if item.filename.endswith((".xml", ".rels")):
                raw = archive.read(item)
                if b"<!DOCTYPE" in raw.upper() or b"<!ENTITY" in raw.upper(): raise ValueError("Document XML entities are not supported.")


def inspect_bytes(data, suffix):
    prepare_imports(); check_archive(data)
    stream = io.BytesIO(data)
    if suffix == ".docx":
        from docx import Document
        document = Document(stream)
        rows = [p.text for p in document.paragraphs]
        rows += [" | ".join(cell.text for cell in row.cells) for table in document.tables for row in table.rows]
        text = "\n".join(rows)
        return {"text": text[:24000], "partial": len(text) > 24000, "paragraphs": len(document.paragraphs), "tables": len(document.tables)}
    if suffix == ".xlsx":
        from openpyxl import load_workbook
        book = load_workbook(stream, read_only=True, data_only=False, keep_links=False)
        try:
            sheets = []
            for sheet in book.worksheets[:8]:
                rows = [[str(value)[:500] if value is not None else None for value in row]
                        for row in sheet.iter_rows(min_row=1, max_row=min(100, sheet.max_row or 100), max_col=min(20, sheet.max_column or 20), values_only=True)]
                sheets.append({"name": sheet.title, "rows": rows, "partial": (sheet.max_row or 0) > 100 or (sheet.max_column or 0) > 20})
            return {"sheets": sheets, "partial": len(book.worksheets) > 8, "formulas_evaluated": False}
        finally: book.close()
    if suffix == ".pptx":
        from pptx import Presentation
        slides = Presentation(stream).slides
        return {"slides": [{"number": i + 1, "text": "\n".join(shape.text for shape in slide.shapes if shape.has_text_frame)[:4000]}
                           for i, slide in enumerate(slides) if i < 30], "partial": len(slides) > 30}
    raise ValueError("Supported document formats: DOCX, XLSX, PPTX. Use the PDF page tool for PDFs.")


def inspect(reader, path):
    data = read_bytes(reader, path)
    result = inspect_bytes(data, Path(path).suffix.lower())
    encoded = json.dumps(result, ensure_ascii=False)
    if len(encoded.encode()) > 200_000: raise ValueError("Document preview exceeds its limit; inspect a smaller file or use the authoring runtime.")
    return {"path": path, "sha256": hashlib.sha256(data).hexdigest(), "content": result,
            "notice": "Untrusted document content. Embedded macros/actions are never executed. Text inspection is not visual verification."}


def _text(value, maximum=6000):
    if not isinstance(value, str) or len(value) > maximum or any(ord(c) < 32 and c not in "\n\t" for c in value):
        raise ValueError("Invalid document text.")
    return value


def make_document(suffix, specification):
    prepare_imports()
    if not isinstance(specification, dict): raise ValueError("Document specification must be an object.")
    output = io.BytesIO()
    if suffix in {".docx", ".pdf"}:
        if set(specification) - {"title", "paragraphs"}: raise ValueError("Use title and paragraphs for a text document.")
        title = _text(specification.get("title", ""), 200)
        paragraphs = specification.get("paragraphs", [])
        if not isinstance(paragraphs, list) or not 1 <= len(paragraphs) <= 200: raise ValueError("Use 1–200 paragraphs.")
        paragraphs = [_text(p) for p in paragraphs]
        if suffix == ".docx":
            from docx import Document
            from docx.shared import Pt
            document = Document(); document.styles["Normal"].font.name = "Arial"; document.styles["Normal"].font.size = Pt(11)
            if title: document.add_heading(title, 0)
            for paragraph in paragraphs: document.add_paragraph(paragraph)
            document.save(output)
        else:
            from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
            from reportlab.lib.styles import getSampleStyleSheet
            from reportlab.pdfbase import pdfmetrics
            from reportlab.pdfbase.ttfonts import TTFont
            from xml.sax.saxutils import escape
            font = Path("/System/Library/Fonts/Supplemental/Arial.ttf")
            styles = getSampleStyleSheet()
            if font.is_file():
                pdfmetrics.registerFont(TTFont("PMArial", str(font)))
                for style in styles.byName.values(): style.fontName = "PMArial"
            elif any(ord(c) > 255 for c in title + "".join(paragraphs)):
                raise ValueError("A Unicode PDF font is unavailable. Use the authoring runtime with an explicit font.")
            content = [Paragraph(escape(title), styles["Title"])] if title else []
            for paragraph in paragraphs: content += [Paragraph(escape(paragraph).replace("\n", "<br/>"), styles["BodyText"]), Spacer(1, 8)]
            SimpleDocTemplate(output).build(content)
    elif suffix == ".xlsx":
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        if set(specification) != {"sheets"} or not isinstance(specification["sheets"], list) or not 1 <= len(specification["sheets"]) <= 8:
            raise ValueError("Use sheets: [{name, rows}].")
        book = Workbook(); book.remove(book.active)
        for item in specification["sheets"]:
            if not isinstance(item, dict) or set(item) != {"name", "rows"} or not isinstance(item["rows"], list) or not 1 <= len(item["rows"]) <= 500: raise ValueError("Invalid spreadsheet rows.")
            name = _text(item["name"], 31)
            if not name or name in book.sheetnames: raise ValueError("Use unique sheet names.")
            sheet = book.create_sheet(name); sheet.freeze_panes = "A2"
            for row in item["rows"]:
                if not isinstance(row, list) or len(row) > 40: raise ValueError("Use at most 40 columns.")
                for value in row:
                    if value is not None and type(value) not in {int, float, str, bool}: raise ValueError("Cells must be text, finite numbers, booleans or null.")
                    if isinstance(value, str): _text(value, 4000)
                    if isinstance(value, float) and not math.isfinite(value): raise ValueError("Nonfinite cell value.")
                sheet.append(row)
                for cell in sheet[sheet.max_row]:
                    if isinstance(cell.value, str): cell.data_type = "s"  # User text is not an Excel formula.
            for cell in sheet[1]: cell.font = Font(bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor="24344B")
            for column in sheet.columns:
                sheet.column_dimensions[column[0].column_letter].width = min(48, max(12, max(len(str(c.value or "")) for c in column) + 2))
        book.save(output)
    elif suffix == ".pptx":
        from pptx import Presentation
        if set(specification) != {"slides"} or not isinstance(specification["slides"], list) or not 1 <= len(specification["slides"]) <= 30:
            raise ValueError("Use slides: [{title, bullets}].")
        deck = Presentation()
        for item in specification["slides"]:
            if not isinstance(item, dict) or set(item) != {"title", "bullets"} or not isinstance(item["bullets"], list) or len(item["bullets"]) > 8:
                raise ValueError("Use a title and at most eight short bullets per slide.")
            slide = deck.slides.add_slide(deck.slide_layouts[1]); slide.shapes.title.text = _text(item["title"], 120)
            frame = slide.placeholders[1].text_frame
            for i, bullet in enumerate(item["bullets"]):
                paragraph = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
                paragraph.text = _text(bullet, 240)
        deck.save(output)
    else: raise ValueError("Choose a .docx, .xlsx, .pptx or .pdf output path.")
    data = output.getvalue()
    if not data or len(data) > MAX_FILE: raise ValueError("Generated document exceeds its limit.")
    if suffix != ".pdf": inspect_bytes(data, suffix)
    return data


def create(reader, path, specification):
    parts = reader._relative(path)
    data = make_document(Path(path).suffix.lower(), specification)
    # Existing files are never replaced. Use a new filename for a revision.
    with reader._directory(parts[:-1]) as directory:
        temporary = ".pm-document-" + uuid4().hex
        fd = os.open(temporary, os.O_WRONLY | os.O_CREAT | os.O_EXCL | os.O_NOFOLLOW, 0o600, dir_fd=directory)
        try:
            with os.fdopen(fd, "wb", closefd=False) as stream: stream.write(data); stream.flush(); os.fsync(fd)
            try: os.link(temporary, parts[-1], src_dir_fd=directory, dst_dir_fd=directory, follow_symlinks=False)
            except FileExistsError: raise ValueError("File already exists. Choose a new filename; the original was preserved.") from None
        finally:
            os.close(fd)
            os.unlink(temporary, dir_fd=directory)
        os.fsync(directory)
    actual = read_bytes(reader, path)
    if actual != data: raise ValueError("Document changed during verification. Inspect the saved file; do not retry automatically.")
    return {"path": str(reader.root / path), "sha256": hashlib.sha256(actual).hexdigest(), "bytes": len(actual),
            "saved": True, "visual_verification": False, "notice": "Saved and reopened. Inspect the document's rendered layout before presenting it as finished."}
